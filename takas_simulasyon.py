"""Ticari takas taslagini stok-akis simülasyonuna ekler.

Kaynaklar:
  sektor_makro_simulasyon.xlsx  stoklar (22 kod, 2025 bilanço / 2026-8 banka)
  ticari takas kanun taslagi     uc katman: kamu faizi, e-fatura takasi, fazla gun
  hekis_enflasyon.py             ileri patika pi = Agustos 2026 yillik TUFE

Excel hucreleri kaymis: sektor H/I, B7'yi gun ve B3'u enflasyon saniyor.
B7 mevduat orani, B8 90 gun, B4 TUFE, B6 ticari kredi faizidir.
Hesap kanun metnine ve Makro notuna gore kurulur: tek agiz 360,1 milyar,
fazla 90 gun, 3 yil, baz %52, taslak HEKIS enflasyonu. %89 kullanilmaz.
"""
import json
from pathlib import Path

import numpy as np
import openpyxl

from akis_model import IDX, SEC, L_G, load, panel, simulate
from hekis_enflasyon import load_obs, inflation_paths, nfp_2026q1, reel_degisim

XLSX = Path(r"c:\Users\asayd\Downloads\sektor_makro_simulasyon.xlsx")


def stoklar():
    wb = openpyxl.load_workbook(XLSX, data_only=False)
    v = wb["Varsayim"]
    asm = {}
    for row in v.iter_rows(min_row=2, max_col=2):
        if row[0].value:
            asm[row[0].value] = row[1].value
    sec = []
    ws = wb["Sektor"]
    for row in ws.iter_rows(min_row=2, max_col=6):
        if row[0].value is None:
            continue
        borc, alacak = row[1].value, row[2].value
        sec.append({
            "kod": row[0].value,
            "borc": borc,
            "alacak": alacak,
            "net": borc - alacak,
            "kasa": row[4].value,
            "karsilik": row[5].value,
        })
    return asm, sec


def bedel(anapara, gun, oran, yil):
    """Fazla gunun yillik bedeli x yil. Anapara milyar TL."""
    return anapara * (gun / 365.0) * oran * yil


def main():
    asm, sec = stoklar()
    pi = inflation_paths(load_obs())["hold_last"][0]
    r_kredi = asm["Ticari kredi faizi"]
    gun = asm["Fazla vade"]
    yil = 3
    borc = asm["Ticari borç, dosya"]
    alacak = asm["Ticari alacak, dosya"]
    net = asm["Net açık"]
    karsilik = asm["Şüpheli karşılık"]
    # Madde 2: karsilik ayrilmis alacak takasa girmez.
    kapanabilir = alacak - karsilik
    kalan_brut_tavan = borc - kapanabilir
    baz = bedel(net, gun, r_kredi, yil)
    taslak = bedel(net, gun, pi, yil)
    fark = baz - taslak
    baz_yil = bedel(net, gun, r_kredi, 1)
    taslak_yil = bedel(net, gun, pi, 1)

    stab, _, tcs, _ = load()
    A, Y = panel(stab, tcs, "2026-Q1")
    L, G = L_G(A, Y)
    nfp = nfp_2026q1()
    # Fark: banka ticari kredi spreadini tutamaz (madde 7/3); net borclu firma odenmez.
    # Tutar bin TL. 1 milyar TL = 1e6 bin TL.
    x = np.zeros(5)
    x[IDX["F"]] = fark * 1e6
    x[IDX["B"]] = -fark * 1e6
    sim = {}
    for lam in (0.0, 0.5, 0.8):
        etki = simulate(L, G, nfp, x, lam=lam) / 1e6
        sim[f"lam{lam}"] = dict(zip(SEC, etki.round(3).tolist()))

    borclular = sorted((s for s in sec if s["net"] > 0), key=lambda s: -s["net"])
    out = {
        "pi": pi,
        "r_kredi": r_kredi,
        "r_mevduat": asm["Vadeli mevduat"],
        "r_politika": asm["Politika faizi, Ekim 2026"],
        "fazla_gun": gun,
        "ticari_borc": borc,
        "ticari_alacak": alacak,
        "net_acik": net,
        "karsilik": karsilik,
        "kapanabilir_alacak": round(kapanabilir, 1),
        "kalan_brut_borc_tavan": round(kalan_brut_tavan, 1),
        "kasa_acigi": asm["Kasa açığı"],
        "banka_kredisi": asm["Banka kredisi 2026/8"],
        "takip": asm["Takip 2026/8"],
        "borc_ozkaynak": asm["Borç / öz kaynak 2025"],
        "gsyh": asm["GSYH 2025"],
        "borc_gsyh_baz": round(borc / asm["GSYH 2025"], 4),
        "borc_gsyh_tavan": round(kalan_brut_tavan / asm["GSYH 2025"], 4),
        "fazla_gun_3y_baz_milyar": round(baz, 3),
        "fazla_gun_3y_taslak_milyar": round(taslak, 3),
        "fark_3y_milyar": round(fark, 3),
        "fark_1y_milyar": round(baz_yil - taslak_yil, 3),
        "reel_1yil_trilyon": dict(zip(SEC, (reel_degisim(nfp, pi) / 1e9).round(3).tolist())),
        "dogrudan_3y_milyar": {"F": round(fark, 3), "B": round(-fark, 3)},
        "sim": sim,
        "net_borclu": [
            {"kod": s["kod"], "net": round(s["net"], 1), "karsilik": round(s["karsilik"], 4)}
            for s in borclular
        ],
    }
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
