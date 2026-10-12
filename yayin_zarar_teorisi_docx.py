"""Zarar Teorisi SFC — bulgu, çözüm yöntemi, mekanizma, sınır: DOCX yayın.

Çalıştır: python3 yayin_zarar_teorisi_docx.py
Çıktı: ZARAR_TEORISI_MINSKY_BULGULAR.docx
"""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

KOK = Path(__file__).resolve().parent
OUT = KOK / "ZARAR_TEORISI_MINSKY_BULGULAR.docx"


def _load(name: str) -> dict:
    p = KOK / name
    if p.exists():
        return json.loads(p.read_text(encoding="utf-8"))
    return {}


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            t.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def build() -> Document:
    uc = _load("uc_kural_senaryo_sonuc.json")
    pon = _load("ponzi_egilim_sonuc.json")
    av = _load("acik_veri_ozet.json")
    ht = _load("hane_taksit_ozeti.json")
    lam = _load("kalibre_lambda_ozeti.json")
    gt = _load("geri_test_ozeti.json")
    dis = _load("akis_model_ozeti.json").get("dis_borc_artisi", {})
    k2031 = uc.get("meta", {}).get("karsilastirma_2031Q4", {})

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Zarar Teorisi Stok-Akış Modeli\n")
    r.bold = True
    r.font.size = Pt(16)
    r2 = title.add_run(
        "Bulgular, çözüm yöntemi, mekanizma ve sınırlar\n"
        "(Minsky rejimi — kur, borç, yaşam maliyeti odaklı okuma)\n\n"
    )
    r2.font.size = Pt(12)
    r3 = title.add_run(f"Depo yayın metni — {date.today().isoformat()}\n")
    r3.italic = True
    r3.font.size = Pt(10)
    doc.add_paragraph(
        "Bu belge politika reçetesi değildir. Ölçüm, varsayım ve senaryo ayrımı korunur. "
        "Sayılar EVDS/BDDK köprüsü (7 Ekim 2026 çekimi) ve senaryo koşusuna bağlıdır; "
        "yenilemek için OKU.txt çalıştırma sırasına bakın."
    )

    doc.add_heading("Özet", level=1)
    doc.add_paragraph(
        "Türkiye finansal hesapları (TCMB EVDS) üzerine kurulu kapalı muhasebe iskeleti: "
        "NFP_t = NFP_{t−1} + NL_t (işlem) + R_t (değerleme). Ana muhasebe okuması: değerleme kayıpları "
        "(özellikle F) banka provizyonuna otomatik yansımaz; yük gecikmeli tanınır ve birikebilir. "
        "Temerrüt sıfır toplamlı transfer değildir (bireysel iflas borcu silmez varsayımı — doğrulanmadı). "
        "Minsky moment için takvim üretilmez; yalnızca kur + borç stoku + yaşam maliyeti kanallarıyla "
        "roll-over riski ve kopuş koşulları tartışılır."
    )

    doc.add_heading("1. Çözüm yöntemi (analitik çerçeve)", level=1)
    doc.add_heading("1.1 Stok-akış tutarlılık", level=2)
    doc.add_paragraph(
        "Her hazne için stok değişimi işlem ve değerleme ile kapanır (2021’den beri işlem akışlarında "
        "panel kapanışı raporlanır). Çözüm, bilanço kimliğini sağlamak ve şokları NL/R üzerinden "
        "haznelere yazmaktır; davranışsal optimum veya genel denge çözücüsü değildir."
    )
    doc.add_heading("1.2 Veri sınıfları", level=2)
    _add_table(
        doc,
        ["Sınıf", "Anlam", "Örnek"],
        [
            ["Ölçüm", "EVDS/BDDK/TÜİK serisi", "Brüt dış borç USD, 13 ay stok büyümesi, TÜFE"],
            ["Varsayım", "Modele elle konmuş pay/eşik", "Ha borç payı D_A_BASE=%40, L dağıtımı"],
            ["Senaryo", "Karşı-olgu yolu", "uc_kural_kilitli, maliyet_ucuz_revize"],
            ["Doğrulanmadı", "Iddia açık", "Bireysel iflas, altın/SDR uyumsuzluk nedeni"],
        ],
    )
    doc.add_heading("1.3 Güvenilir yorum kanalları (kısıt)", level=2)
    doc.add_paragraph(
        "Bu yayında ağır basan üç kanal: (1) kur — USD/TRY ve brüt dış borç ayrıştırması, D değerleme; "
        "(2) borç kalemleri — hane/firma/kamu/dış stok ve büyüme hızları; "
        "(3) yaşam maliyeti — enflasyon, taban/sepet, taksit/gelir (tanımlı nesneler). "
        "TGA, VYS net etkisi, lambda, desil dağılımı sonuç satırı gibi okunmaz."
    )
    doc.add_heading("1.4 Test ve kalibrasyon", level=2)
    t2 = gt.get("T2_matris_testi", {})
    r2_lm = t2.get("R2_Lm_modeli(havuz)", "0,868")
    r2_basit = t2.get("R2_kiyas_toplam_varlik_payi", "0,862")
    doc.add_paragraph(
        f"Geriye dönük Lm testi: orantılı Lm R² ≈ {r2_lm}; basit varlık payı ≈ {r2_basit} — matris ek bilgi "
        "eklemez. Lambda (F değerleme → provizyon): GSYH normalize regresyon anlamsız/negatif; "
        "geçiş üst sınır ~%1,4 (kalibre_lambda_ozeti.json). F→K ~%24 bulgusu geri çekildi (ortak trend)."
    )

    doc.add_heading("2. Mekanizma", level=1)
    doc.add_heading("2.1 Hazneler ve akışlar", level=2)
    doc.add_paragraph(
        "H (hane S.14+15), F (firmalar), B′ (bankalar − TCMB), K′ (Kamu+TCMB), D (dış), U (uyumsuzluk). "
        "Bütçe kanalı K′ için merkezi yönetim vergi/faiz/harcama yazar; mükellef–hazne ayrımı varsayımdır."
    )
    doc.add_heading("2.2 Zarar ve tanıma (yük stoku)", level=2)
    doc.add_paragraph(
        "F’nin net değerleme zararı ile bankanın gerçekleşen kredi zararı (provizyon) serisi arasında "
        "ampirik geçiş zayıftır. Muhasebe yorumu: zarar F’de kalabilir veya forbearance/yapılandırma ile "
        "gecikmeli tanınır — kurumsal motor modellenmez. Temerrüt çift kayıt: alacaklı karşılık ayırır; "
        "borçlu defterinde borç nominal kalıp gecikme faizi birikebilir (faiz oranı varsayım)."
    )
    doc.add_heading("2.3 L / Lm — karşı taraf dağıtımı", level=2)
    doc.add_paragraph(
        "Borçlu haznenin değerleme/işlem kaybı, alacaklı varlık paylarına bölünür (Lm). Ölçülmüş fatura "
        "ağı değildir; sözleşmesel dağıtım. Banka serbest sermayeyi aşan kısım senaryolarda L[B,·] ile "
        "yayılır; fiilen çoğunlukla kamu desteği kanalıyla okunur (varsayım)."
    )
    doc.add_heading("2.4 Üç kanalın birleşik mekanizması → Minsky okuması", level=2)
    doc.add_paragraph(
        "Kur ↑ → TL cinsinden borç ağırlığı ve D tarafında değerleme (net pozisyon ≠ brüt dış borç). "
        "Borç stoku hızlı büyür (13 ay kart/tüketici ~%43–46) → roll-over / yeni borçla eskiyi ödeme "
        "davranışı. Yaşam maliyeti ↑ → reel taban erir; servis/gelir kötüleşir. "
        "Minsky moment (teori): finansmanın kesilmesi + gelir/kur şoku; modelde proxy eşikler "
        "hedge ≤%30, ponzi etiketi >%40 (Ha servis/gelir, senaryo parametresi)."
    )
    doc.add_paragraph(
        "Hedge: gelir faiz+anaparayı karşılar. Spekülatif: faiz taşınır, anapara yeniden borçlanır. "
        "Ponzi (davranış): gelir faizi karşılamaz; yeni borç eskiyi besler. Kopuş: kart/limit/kredi "
        "büyümesinin durması — depoda spread/red/limit serisi yok; tarih atanmaz."
    )
    doc.add_heading("2.5 uc_kural senaryo geri bildirimi", level=2)
    doc.add_paragraph(
        "Taban, kapsama, DTI, tüketim tavanı, limit_dondur, F ücret kanalı, K′ bütçe NL birlikte çalışır. "
        "Revize: reel yaşam maliyeti düşüşü + makas; ek kilit limit_dondur = limit_stres ∧ yaşam düşüş eşiği. "
        "Sert: agresif taban, yüksek F ücret yükü, sürekli limit_dondur senaryoda."
    )

    doc.add_heading("3. Bulgular (ölçüm + senaryo)", level=1)
    doc.add_heading("3.1 Yaşam maliyeti ve taban (ölçüm/köprü)", level=2)
    tg = av.get("taban_girdi", {}) or uc.get("meta", {}).get("taban_girdi", {})
    tufe = tg.get("tufe_yillik_2026") or av.get("olcumler", {}).get("tufe_yillik", {}).get("deger", "—")
    doc.add_paragraph(
        f"TÜFE yıllık (köprü): ~{tufe}%. Asgari net ~28.076 TL/ay; bekar yaşam sepeti ~48.885 TL/ay → "
        f"kapsama ~%57,4 (2026 giriş). hane_taksit (2024, kart hariç): taksit/gelir ~%4,0 (18 yıl vade varsayımı); "
        f"faiz stok oranı ~%53 — oran yükselişi çoğunlukla faiz faturası."
    )
    doc.add_heading("3.2 Borç kalemleri (ölçüm)", level=2)
    kart = av.get("taban_girdi", {}).get("kart_buyume_13h_yillik") or tg.get("kart_buyume_13h_yillik", "—")
    tuk = av.get("taban_girdi", {}).get("tuketici_kredi_13h_yillik") or tg.get("tuketici_kredi_13h_yillik", "—")
    doc.add_paragraph(
        f"BDDK/EVDS 13 ay stok büyümesi: kart ~%{kart}, tüketici ~%{tuk}. "
        "Dar takip G1 hane ~%5,09 (2026-09). Geniş tanım (yakın izleme) modele girmez; OKU ~%20–28 bandı."
    )
    doc.add_heading("3.3 Kur ve dış borç (ölçüm)", level=2)
    ay = dis.get("ayristirma_2023Q4_to_son", {}) or {}
    top = ay.get("Toplam", {})
    doc.add_paragraph(
        f"Brüt dış borç 2023-Q4 → 2026-Q2: EVDS ayrıştırma (Toplam) — yeni borçlanma TL ~{top.get('yeni_borclanma_TL_trilyon', '—')} trilyon, "
        f"kur etkisi ~{top.get('kur_etkisi_TL_trilyon', '—')} trilyon (OKU: kur payı ~%75). "
        "D net NFP değişimi = net işlem + değerleme; brüt USD stoku ile aynı nesne değil."
    )
    doc.add_heading("3.4 Ha servis proxy ve Minsky bandı (senaryo)", level=2)
    rows_pon = []
    oz = pon.get("ozet", {})
    for sen in ("mevcut", "uc_kural_kilitli", "maliyet_ucuz_revize"):
        o = oz.get(sen, {})
        rows_pon.append(
            [
                sen,
                str(o.get("servis_son", "")),
                str(o.get("ponzi_fazlasi_pp_son", "")),
                str(o.get("yon", "")),
            ]
        )
    _add_table(doc, ["Senaryo", "Ha servis/gelir (2031-Q4)", "Ponzi fazlası (pp)", "Eğilim 2026–2031"], rows_pon)
    doc.add_paragraph(
        "Model Ha servis ~%47 (2026-Q4 mevcut giriş) ile hane_taksit ~%4 aynı nesne değildir; birlikte okunmalıdır."
    )

    doc.add_heading("3.5 uc_kural 2031-Q4 karşılaştırma (senaryo)", level=2)
    rows_uc = []
    keys = [
        ("borc_servis_orani_Ha", "Ha servis/gelir"),
        ("kapsama", "Kapsama"),
        ("yasam_maliyeti_TL", "Yaşam maliyeti (TL/ay)"),
        ("net_taban_TL", "Net taban (TL/ay)"),
        ("NFP_H", "NFP_H (milyar TL)"),
        ("F_ucret_kanali_kumul_milyar", "F ücret kanalı küm."),
        ("korunmasiz_doviz_stok_mr_usd", "Korunmasız USD stok (mr)"),
    ]
    for key, label in keys:
        block = k2031.get(key, {})
        rows_uc.append(
            [
                label,
                str(block.get("mevcut", "")),
                str(block.get("uc_kural_kilitli", "")),
                str(block.get("maliyet_ucuz_revize", "")),
            ]
        )
    _add_table(doc, ["Gösterge", "mevcut", "uc_kural_kilitli", "maliyet_ucuz_revize"], rows_uc)

    doc.add_heading("3.6 Hane alt taban (varsayım)", level=2)
    ha = _load("akis_model_ozeti.json").get("hane_alt", {})
    tab = ha.get("taban_NFP_trilyonTL", {}).get("Ha asgari", {})
    doc.add_paragraph(
        f"D_A_BASE=%40 (taban; ölçüm değil). 2026-Q1 Ha NFP ~{tab.get('NFP', '−1,63')} trilyon TL. "
        "Duyarlılık tablosunda satir_tipi: taban | duyarlilik (hane_alt_duyarlik.csv). "
        "%70 satırı duyarlılıktır; lambda_B taban ~0,16, %70 duyarlılık ~0,59."
    )

    doc.add_heading("4. Minsky noktası — yaklaşma ve bilinemeyenler", level=1)
    doc.add_paragraph(
        "Tanım: nakit akışı bağının kopması ve/veya roll-over’ın durması (yeni kredi/kart hattı kesilir). "
        "Ölçüde henüz sistemik kredi durması ve spread/limit serisi yok → moment atanmaz. "
        "Yaklaşma imzası: yüksek stok büyümesi + yüksek faiz ortamı + geçim baskısı + kurun TL yükünü artırması. "
        "Sert senaryo servisi düşürür ve limit_dondur uygular — karşı-olgu, tahmin değil."
    )
    _add_table(
        doc,
        ["Senaryo çerçevesi", "Yön (Minsky okuması)"],
        [
            ["Kur ↑ + yaşam ↑", "Servis baskısı artar; roll-over ihtiyacı artar"],
            ["Stok büyümesi → 0", "Moment tanımına en yakın olay (Ponzi beslemesi kesilir)"],
            ["Yaşam ↓, stok aynı", "Geçim rahatlar; borç yolu Ponzi kalırsa moment ertelenir"],
            ["Kur ↑ + stok hızlı", "Üçlü şok; kopuş riski artar (büyüklük desilde bilinmez)"],
        ],
    )

    doc.add_heading("5. Model sınırları", level=1)
    doc.add_paragraph(
        "Merkezi sınır: kapalı muhasebe her kaydı bir yere yazar; tanımanın neden/ne zaman/nasıl olduğunu "
        "açıklamaz (MODEL_KUSURLARI.md). Lambda zayıf; Lm basit paya yakın; hisse/FDI/YP kırılımı yok; "
        "K′ ≠ merkezi bütçe; Minsky psikolojisi/forbearance yok. Politika reçetesi üretilmez."
    )
    doc.add_heading("5.1 Ne denmez", level=2)
    for line in [
        "“Türkiye Minsky moment’i 20XX’te geçti.”",
        "“Ha servis %47 = her hane taksit/gelir %47.”",
        "“Lm matrisi gerçek borç ağını ölçtü.”",
        "“Lambda_B = 0,59 kesin tanıma oranıdır.” (%70 duyarlılık satırı.)",
        "Varsayım paylarını ölçülmüş gerçek gibi sunmak.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("6. Dosya referansları ve yenileme", level=1)
    refs = [
        "OKU.txt — veri, test, lambda, hane_alt",
        "MODEL_KUSURLARI.md, VARSAYIM_UYARISI.md",
        "MINSKY_GUNCEL_KARSILASTIRMA.md",
        "uc_kural_senaryo_sonuc.json, ponzi_egilim_sonuc.json",
        "acik_veri_ozet.json, hane_taksit_ozeti.json, kalibre_lambda_ozeti.json",
        "Çalıştırma: python3 akis_model.py; python3 uc_kural_senaryo.py; python3 yayin_zarar_teorisi_docx.py",
    ]
    for r in refs:
        doc.add_paragraph(r, style="List Bullet")

    doc.add_paragraph(
        "\n— Belge otomatik üretilmiştir; markdown kaynaklarla çelişirse depo JSON/CSV ve OKU.txt esas alınır."
    )
    return doc


def main():
    doc = build()
    doc.save(OUT)
    print(f"Yazildi: {OUT}")


if __name__ == "__main__":
    main()
